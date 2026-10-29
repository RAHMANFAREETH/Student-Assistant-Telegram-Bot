"""V9.0 Academic Toolkit: bounded, local first student utilities.

Text extraction reuses V8. No model claims, hidden network calls or source execution.
Optional third-party OCR/metadata calls live in v9_bot.py behind explicit requests.
"""
from __future__ import annotations
import ast
import csv
import difflib
from collections import Counter
from datetime import datetime
from io import BytesIO, StringIO
import math
import re
from typing import Iterable

from v8_engine import extract_document, StudyInputError, tokens

MAX_UPLOAD = 8 * 1024 * 1024


def all_text(pages):
    return '\n'.join(s for _, s in pages)[:90000]


def read_text(filename, raw):
    return all_text(extract_document(filename, raw))


def assignment_report(filename, raw, min_words=500, max_words=3000):
    text = read_text(filename, raw)
    words = re.findall(r"\b[\w’-]+\b", text, re.UNICODE)
    size = len(words)
    headings = {h: bool(re.search(r'(?im)^\s*' + re.escape(h) + r'\s*:?\s*$', text))
                for h in ('Introduction', 'Conclusion', 'References')}
    has_inline = bool(re.search(r'\([A-Z][\w-]+,?\s+\d{4}[a-z]?\)|\[\d{1,3}\]', text))
    notes = [f'📋 Assignment checklist: {filename}', f'Words (approx.): {size}',
             f'Target range: {min_words}–{max_words}']
    notes.append('Word limit: ' + ('✅ In range' if min_words <= size <= max_words else '⚠️ Outside target'))
    notes.extend(f'{"✅" if ok else "⚠️"} {k} heading detected' for k, ok in headings.items())
    notes.append(('✅' if has_inline else '⚠️') + ' Common in-text citation pattern detected')
    notes.append('Check the actual assignment rubric, plagiarism rules and formatting manually. This is a heuristic check.')
    return '\n'.join(notes)


def syllabus_lines(text):
    lines = []
    for row in text.splitlines():
        x = re.sub(r'^\s*(?:[\d]+[).:-]|[-*•])\s*', '', row).strip(' \t:')
        if 3 <= len(x) <= 120 and x.casefold() not in {r.casefold() for r in lines}:
            lines.append(x)
        if len(lines) >= 60:
            break
    if not lines:
        raise StudyInputError('Send at least one topic per line (3–120 characters).')
    return lines


def formula_lines(text):
    found = []
    for raw in text.splitlines():
        line = ' '.join(raw.strip().split())
        if 3 <= len(line) <= 170 and ('=' in line or re.search(r'\b(?:sin|cos|tan|log|sqrt|integral)\b|[∑∫√≈≤≥]', line, re.I)):
            if line not in found:
                found.append(line)
        if len(found) >= 60:
            break
    return found


def formula_docx(lines):
    from docx import Document
    doc = Document()
    doc.add_heading('Study Formula Sheet', 0)
    for line in lines:
        doc.add_paragraph(line, style='List Bullet')
    doc.add_paragraph('Verify every formula against your source; extraction may include ordinary equations.')
    output = BytesIO()
    doc.save(output); output.seek(0); output.name = 'formula_sheet.docx'
    return output


def prepare_deck(text):
    """A portable, editable Anki TSV, not an .apkg package."""
    rows = [('Front', 'Back')]
    for part in text:
        q, a = str(part[0])[:500], str(part[1])[:1800]
        q = q.replace('\r', ' ').replace('\n', '<br>').replace('\t', ' ')
        a = a.replace('\r', ' ').replace('\n', '<br>').replace('\t', ' ')
        # Prevent spreadsheet formula injection on opening TSV in Excel.
        q = "'" + q if q.lstrip().startswith(('=', '+', '-', '@')) else q
        a = "'" + a if a.lstrip().startswith(('=', '+', '-', '@')) else a
        rows.append((q, a))
    out = StringIO(newline='')
    csv.writer(out, delimiter='\t', lineterminator='\n').writerows(rows)
    data = BytesIO(out.getvalue().encode('utf-8-sig'))
    data.name = 'anki_import.tsv'
    return data


def document_diff(text_a, text_b):
    a, b = text_a.splitlines(), text_b.splitlines()
    diff = list(difflib.unified_diff(a[:1200], b[:1200], fromfile='First', tofile='Second', lineterm='', n=1))
    if not diff:
        return '✅ No line-level changes found in selectable text.'
    additions = sum(x.startswith('+') and not x.startswith('+++') for x in diff)
    removals = sum(x.startswith('-') and not x.startswith('---') for x in diff)
    excerpt = '\n'.join(diff[2:][:55])
    return f'🔎 Text differences: {additions} added and {removals} removed lines (approx.).\n\n{excerpt[:2900]}\n\nImages and layout differences are not compared.'


def mindmap_topics(text):
    candidates=[]
    for line in text.splitlines():
        s = re.sub(r'^\s*[0-9]+[.):]?\s*|^\s*[-*•]\s*', '', line).strip()
        if 4 <= len(s) <= 75 and len(s.split()) <= 12 and s not in candidates:
            candidates.append(s)
        if len(candidates) >= 10: break
    if len(candidates) < 3:
        counts = Counter(tokens(text))
        for word, _ in counts.most_common(10):
            if word not in candidates:
                candidates.append(word.title())
    return candidates[:10]


def mindmap_png(title, topics):
    from PIL import Image, ImageDraw, ImageFont
    w, h = 1100, 160 + 92 * max(len(topics), 2)
    image = Image.new('RGB', (w, h), '#f5f7fb')
    draw = ImageDraw.Draw(image)
    fonts = ['/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf',
             '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
             'C:/Windows/Fonts/arial.ttf']
    font = next((ImageFont.truetype(p, 24) for p in fonts if _file_exists(p)), ImageFont.load_default())
    bold = next((ImageFont.truetype(p, 29) for p in fonts if _file_exists(p)), ImageFont.load_default())
    draw.rounded_rectangle((25, 22, w-25, 112), radius=22, fill='#183657')
    draw.text((49, 49), title[:42], fill='white', font=bold)
    for i, topic in enumerate(topics):
        y=145+i*92
        draw.line((150, y+20, 150, y+51), fill='#7188ab', width=3)
        draw.line((150, y+36, 182, y+36), fill='#7188ab', width=3)
        draw.rounded_rectangle((180, y, w-45, y+65), radius=14, fill='#dceafc')
        draw.text((201, y+18), topic[:65], fill='#183657', font=font)
    out=BytesIO(); image.save(out, format='PNG'); out.seek(0); out.name='study_mindmap.png'; image.close()
    return out


def _file_exists(p):
    from os.path import isfile
    return isfile(p)


def presentation_pptx(text, title='Study Presentation'):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    lines = [re.sub(r'^\s*(?:[-*•]|\d+[.)])\s*', '', x).strip() for x in text.splitlines()]
    lines=[x for x in lines if x and len(x)>2][:90]
    if not lines: raise StudyInputError('Your notes contain no usable text.')
    prs=Presentation(); prs.slide_width= Inches(13.333); prs.slide_height= Inches(7.5)
    cover=prs.slides.add_slide(prs.slide_layouts[0]); cover.shapes.title.text=title[:75]
    cover.placeholders[1].text='Editable slides generated from your notes. Review for accuracy.'
    for index in range(0, min(len(lines), 72), 8):
        block=lines[index:index+8]
        slide=prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text=f'Key points {index//8 + 1}'
        frame=slide.placeholders[1].text_frame; frame.clear()
        for i, sentence in enumerate(block):
            p=frame.paragraphs[0] if i==0 else frame.add_paragraph()
            p.text=sentence[:210]; p.level=0; p.font.size=Pt(22)
        if len(prs.slides)>=11: break
    out=BytesIO(); prs.save(out); out.seek(0); out.name='study_slides.pptx'
    return out, len(prs.slides)


# Safe symbolic arithmetic. Never sympify/parse_expr/eval raw student input.
def math_solution(input_text):
    from sympy import Symbol, Eq, solve, expand, Poly, pretty, Integer, Rational
    raw = input_text.strip().replace('²','^2').replace('×','*').replace('÷','/')
    raw = raw.replace('^','**').replace('−','-')
    if not 3 <= len(raw) <= 105 or '=' not in raw:
        raise StudyInputError('Enter a simple equation such as 2*x + 3 = 11 or x^2 - 4 = 0.')
    if not re.fullmatch(r'[xX0-9. +*/()=\-]+', raw):
        raise StudyInputError('Math supports numeric equations in x using +, -, *, / and ^ only.')
    left, *rest=raw.split('=')
    if len(rest)!=1: raise StudyInputError('Use one equals sign.')
    x = Symbol('x')
    def visit(node, depth=0):
        if depth>12: raise StudyInputError('Equation is too complex.')
        if isinstance(node, ast.Expression): return visit(node.body,depth+1)
        if isinstance(node, ast.Name) and node.id.lower()=='x': return x
        if isinstance(node, ast.Constant) and type(node.value) in (int,float):
            if abs(node.value)>1e7: raise StudyInputError('Number is too large.')
            return Rational(str(node.value))
        if isinstance(node,ast.UnaryOp) and isinstance(node.op,(ast.USub,ast.UAdd)):
            a=visit(node.operand,depth+1);return -a if isinstance(node.op,ast.USub) else a
        if isinstance(node,ast.BinOp) and isinstance(node.op,(ast.Add,ast.Sub,ast.Mult,ast.Div,ast.Pow)):
            a,b=visit(node.left,depth+1),visit(node.right,depth+1)
            if isinstance(node.op,ast.Add): return a+b
            if isinstance(node.op,ast.Sub): return a-b
            if isinstance(node.op,ast.Mult): return a*b
            if isinstance(node.op,ast.Div):
                if b==0: raise StudyInputError('Division by zero.')
                return a/b
            if not getattr(b,'is_Integer',False) or not 0 <= int(b) <= 2:
                raise StudyInputError('Only powers 0, 1 and 2 are supported.')
            return a**int(b)
        raise StudyInputError('Unsupported expression. Use explicit multiplication such as 2*x.')
    try:
        parsed_left=ast.parse(left.strip(),mode='eval')
        parsed_right=ast.parse(rest[0].strip(),mode='eval')
        nodes=list(ast.walk(parsed_left))+list(ast.walk(parsed_right))
        if len(nodes)>50 or sum(isinstance(n,ast.Pow) for n in nodes)>2:
            raise StudyInputError('Equation is too complex.')
        lhs=visit(parsed_left);rhs=visit(parsed_right)
        eq=expand(lhs-rhs)
        poly=Poly(eq,x)
        if poly.degree()>2: raise StudyInputError('Only linear and quadratic polynomial equations are supported.')
        answers=solve(Eq(lhs,rhs),x)
    except (StudyInputError,): raise
    except Exception as exc: raise StudyInputError('Could not interpret that equation. Try 2*x+3=11.') from exc
    if not answers:
        if eq == 0: return f'Equation: {input_text}\nAll real values of x satisfy this identity (subject to valid denominators).'
        return f'Equation: {input_text}\nNo solutions found in the supported domain.'
    verified=[]
    for val in answers[:2]:
        if val.is_real is False: verified.append(f'x = {val} (complex)')
        else: verified.append(f'x = {val}')
    return ('🧮 Equation: '+input_text[:110]+'\n\nStep 1. Move all terms to one side:\n'
            +str(eq)+' = 0\nStep 2. Solve the resulting '
            +('linear' if poly.degree()==1 else 'quadratic')+' equation:\n'
            +'\n'.join(verified)+'\n\nCheck: substitute each solution into the original equation.')


def parse_timetable(text):
    """Recognizes lines like Mon 09:00-10:30 Chemistry; user must confirm before saving."""
    weekdays={'mon':0,'tue':1,'wed':2,'thu':3,'fri':4,'sat':5,'sun':6}
    rows=[]
    pat = re.compile(
        r'^\s*(Mon(?:day)?|Tue(?:sday)?|Wed(?:nesday)?|Thu(?:rsday)?|Fri(?:day)?|Sat(?:urday)?|Sun(?:day)?)\s+'
        r'(\d{1,2}:\d{2})\s*[-–]\s*(\d{1,2}:\d{2})\s+(.{3,70})\s*$',
        re.I
    )
    for line in text.splitlines():
        m=pat.match(line)
        if not m: continue
        day=weekdays[m[1][:3].lower()]
        try:
            a=datetime.strptime(m[2],'%H:%M').time(); b=datetime.strptime(m[3],'%H:%M').time()
        except ValueError: continue
        if a>=b: continue
        title=m[4].strip()
        if len(rows)<12:rows.append((day, a, b, title))
    return rows


def diagram_labels(text):
    """Text-label reader, NOT a vision diagram interpreter."""
    labels=[x.strip() for x in text.splitlines() if 2 <= len(x.strip()) <= 70][:25]
    if not labels:raise StudyInputError('OCR found no readable labels.')
    return ('🔬 Detected diagram labels (OCR):\n'+ '\n'.join('• '+s for s in labels[:15])+
            '\n\nThis does NOT infer arrows, spatial relationships or scientific accuracy. '
            'Use the original image and verify each label. A vision model is needed for full visual explanations.')
