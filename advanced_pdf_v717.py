"""V7.17 advanced PDF tools for Student Assistant Telegram Bot.

Five opt-in local-only tools using pypdf, PDFium2 and ReportLab. No database, cloud API, or new Telegram bot.
Each job has bounded input/output size and shares one worker slot to limit RAM.
"""
from __future__ import annotations

from html import escape
from io import BytesIO
from pathlib import Path
import re
import threading
import time
from zipfile import ZipFile, ZIP_DEFLATED

MAX_INPUT_BYTES = 10 * 1024 * 1024
MAX_OUTPUT_BYTES = 18 * 1024 * 1024
MAX_PAGES = 50
MAX_IMAGE_PAGES = 20
MAX_WORD_PAGES = 35
MAX_WORD_TEXT_CHARS = 250_000
MAX_OUTPUT_PAGES = 60
MAX_SESSIONS = 10
SESSION_TTL_SECONDS = 15 * 60

COMPRESS_BUTTON = "🗜 Compress PDF"
PDF_WORD_BUTTON = "📄 PDF → Word"
PDF_IMAGES_BUTTON = "🖼 PDF → Images"
ORGANIZE_BUTTON = "🔀 Organize PDF"
WATERMARK_BUTTON = "💧 Add Watermark"
ADVANCED_BUTTONS = (
    COMPRESS_BUTTON, PDF_WORD_BUTTON, PDF_IMAGES_BUTTON, ORGANIZE_BUTTON, WATERMARK_BUTTON,
)
COMPRESS_LEVELS = {"🟢 Light": (None, None, 0), "🟡 Balanced": (180, 120, 75), "🔴 Strong": (140, 95, 55)}
IMAGE_FORMATS = {"🖼 PNG": "png", "📷 JPG": "jpg"}
WATERMARK_PLACEMENTS = {"⬆️ Top": "top", "🎯 Center": "center", "⬇️ Bottom": "bottom"}


class AdvancedPdfError(ValueError):
    """Expected input/processing problem safe to report to Telegram users."""


def _pdf_reader(blob: bytes):
    from pypdf import PdfReader
    try:
        reader = PdfReader(BytesIO(blob), strict=False)
        if reader.is_encrypted:
            raise AdvancedPdfError("Password-protected PDFs aren't supported. Unlock your own file first.")
        return reader
    except AdvancedPdfError:
        raise
    except Exception as exc:
        raise AdvancedPdfError("This PDF appears invalid or damaged. Try another PDF.") from exc


def _out(blob: bytes, name: str) -> BytesIO:
    if not blob or len(blob) > MAX_OUTPUT_BYTES:
        raise AdvancedPdfError("Result is empty or too large for this bot (18 MB limit).")
    output = BytesIO(blob)
    output.name = name
    return output


def safe_stem(filename: str) -> str:
    filename = str(filename or "document").replace("\\", "/").split("/")[-1]
    stem = Path(filename).stem
    stem = re.sub(r"[^\w.\- ]", "_", stem, flags=re.UNICODE).strip(" ._")
    return (stem or "document")[:65]


def pdf_info(blob: bytes) -> int:
    if not isinstance(blob, bytes) or not blob or len(blob) > MAX_INPUT_BYTES:
        raise AdvancedPdfError("PDF must be between 1 byte and 10 MB.")
    reader = _pdf_reader(blob)
    try:
        count = len(reader.pages)
        if count < 1 or count > MAX_PAGES:
            raise AdvancedPdfError(f"Use a PDF containing 1–{MAX_PAGES} pages.")
        return count
    except AdvancedPdfError:
        raise
    except Exception as exc:
        raise AdvancedPdfError("This PDF appears invalid or damaged. Try another PDF.") from exc


def parse_pages(spec: str, total: int, *, max_pages: int = MAX_OUTPUT_PAGES, preserve_duplicates: bool = False) -> list[int]:
    """1-based input; accepts 1,3-5 and descending 5-3; returns zero-based pages."""
    spec = (spec or "").strip()
    if spec.lower() in {"all", "all pages", "keep all", "✅ all pages", "✅ keep all"}:
        selected = list(range(total))
    else:
        if not re.fullmatch(r"[0-9,\-\s]+", spec) or not spec.strip(", -"):
            raise AdvancedPdfError("Use page numbers like 3,1,2-5 or tap All Pages.")
        selected = []
        for part in spec.split(","):
            part = part.strip()
            match = re.fullmatch(r"(\d+)\s*(?:-\s*(\d+))?", part)
            if not match:
                raise AdvancedPdfError("Invalid page list. Example: 3,1,2-5")
            start = int(match.group(1))
            end = int(match.group(2)) if match.group(2) else start
            if not (1 <= start <= total and 1 <= end <= total):
                raise AdvancedPdfError(f"Pages must be between 1 and {total}.")
            step = 1 if end >= start else -1
            if len(selected) + abs(end - start) + 1 > max_pages:
                raise AdvancedPdfError(f"Maximum {max_pages} output pages.")
            selected.extend(n - 1 for n in range(start, end + step, step))
        if not preserve_duplicates:
            selected = list(dict.fromkeys(selected))
    if not selected or len(selected) > max_pages:
        raise AdvancedPdfError(f"Choose 1–{max_pages} pages.")
    return selected


def _write(writer, name: str) -> BytesIO:
    result = BytesIO()
    writer.write(result)
    return _out(result.getvalue(), name)


def compress_pdf(blob: bytes, filename: str, level: str) -> tuple[BytesIO, str]:
    if level not in COMPRESS_LEVELS:
        raise AdvancedPdfError("Choose Light, Balanced or Strong compression.")
    pdf_info(blob)
    from pypdf import PdfWriter
    from PIL import Image
    try:
        writer = PdfWriter()
        writer.append(_pdf_reader(blob))
        dpi_threshold, dpi_target, quality = COMPRESS_LEVELS[level]
        for page in writer.pages:
            if quality:
                for image_file in page.images:
                    try:
                        image = image_file.image
                        if image.width * image.height > 12_000_000:
                            continue
                        # Lossy re-encoding only for nontransparent RGB-style pictures.
                        if image.mode not in ("RGB", "L"):
                            continue
                        if dpi_threshold:
                            x_dpi = image.width * 72 / max(1, float(page.mediabox.width))
                            y_dpi = image.height * 72 / max(1, float(page.mediabox.height))
                            if min(x_dpi, y_dpi) > dpi_threshold:
                                factor = dpi_target / max(x_dpi, y_dpi)
                                if 0 < factor < 1:
                                    image = image.resize((max(1,int(image.width * factor)), max(1,int(image.height * factor))), Image.Resampling.LANCZOS)
                        image_file.replace(image, quality=quality)
                    except (ValueError, TypeError, OSError, NotImplementedError):
                        # Keep unsupported image types intact rather than dropping data.
                        continue
            try:
                page.compress_content_streams()
            except (ValueError, TypeError):
                pass
        writer.compress_identical_objects(remove_identicals=True, remove_orphans=False)
        encoded = BytesIO()
        writer.write(encoded)
        best = encoded.getvalue() if encoded.tell() < len(blob) else blob
    except AdvancedPdfError:
        raise
    except Exception as exc:
        raise AdvancedPdfError("Could not compress this PDF. Try another document.") from exc
    before, after = len(blob)/1024, len(best)/1024
    note = (f"✅ {level} compression: {before:.0f} KB → {after:.0f} KB "
            f"({100*(1-len(best)/len(blob)):.1f}% smaller)." if len(best) < len(blob)
            else "ℹ️ PDF is already optimized; the original size was preserved.")
    if level == "🔴 Strong":
        note += "\n⚠️ Strong mode may reduce embedded image quality; text stays selectable."
    return _out(best, safe_stem(filename) + "_compressed.pdf"), note


def pdf_to_word(blob: bytes, filename: str) -> BytesIO:
    """Text-first Word export. Original PDF graphics and complex layout aren't recreated."""
    pdf_info(blob)
    from docx import Document
    try:
        reader = _pdf_reader(blob)
        if len(reader.pages) > MAX_WORD_PAGES:
            raise AdvancedPdfError(f"PDF → Word accepts up to {MAX_WORD_PAGES} pages.")
        docx = Document()
        extracted = 0
        for i, page in enumerate(reader.pages):
            if i:
                docx.add_page_break()
            text = (page.extract_text(extraction_mode="layout") or "") if "/Contents" in page else ""
            extracted += len(text.strip())
            if extracted > MAX_WORD_TEXT_CHARS:
                raise AdvancedPdfError("This PDF contains too much text for Word export. Split it into smaller parts.")
            for line in text.splitlines():
                docx.add_paragraph(line.rstrip())
        if extracted < 2:
            raise AdvancedPdfError("No selectable text found. Scanned/image-only PDFs require OCR before Word conversion.")
        output = BytesIO(); docx.save(output)
        return _out(output.getvalue(), safe_stem(filename) + ".docx")
    except AdvancedPdfError:
        raise
    except Exception as exc:
        raise AdvancedPdfError("PDF → Word conversion failed for this document.") from exc


def pdf_to_images(blob: bytes, filename: str, fmt: str, pages: list[int]) -> BytesIO:
    pdf_info(blob)
    if fmt not in ("png", "jpg"):
        raise AdvancedPdfError("Choose JPG or PNG.")
    if not pages or len(pages) > MAX_IMAGE_PAGES:
        raise AdvancedPdfError(f"Select 1–{MAX_IMAGE_PAGES} pages.")
    import pypdfium2 as pdfium
    output = BytesIO()
    try:
        with pdfium.PdfDocument(blob) as pdf, ZipFile(output, "w", ZIP_DEFLATED, compresslevel=6) as zf:
            for page_index in pages:
                if page_index < 0 or page_index >= len(pdf):
                    raise AdvancedPdfError("Page selection is outside the PDF.")
                page=pdf.get_page(page_index)
                try:
                    w,h=page.get_size()
                    if w*h*(120/72)**2 > 10_000_000:
                        raise AdvancedPdfError("A page is too large for image export. Try a smaller PDF.")
                    bitmap = page.render(scale=120/72)
                    image = bitmap.to_pil().convert("RGB")
                    image_buf=BytesIO()
                    image.save(image_buf, "PNG" if fmt=="png" else "JPEG", quality=84)
                    zf.writestr(f"page_{page_index + 1:03d}.{fmt}", image_buf.getvalue())
                    image.close()
                    bitmap.close()
                finally:
                    page.close()
                if output.tell() > MAX_OUTPUT_BYTES:
                    raise AdvancedPdfError("Export ZIP is too large. Select fewer pages or JPG.")
        return _out(output.getvalue(), safe_stem(filename) + "_images.zip")
    except AdvancedPdfError:
        raise
    except Exception as exc:
        raise AdvancedPdfError("Could not render PDF pages as images.") from exc


def organize_pdf(blob: bytes, filename: str, pages: list[int]) -> BytesIO:
    pdf_info(blob)
    if not pages or len(pages) > MAX_OUTPUT_PAGES:
        raise AdvancedPdfError("Select at least one page (maximum 60).")
    from pypdf import PdfWriter
    try:
        reader=_pdf_reader(blob)
        writer=PdfWriter()
        for idx in pages:
            if not 0<=idx<len(reader.pages):
                raise AdvancedPdfError("Page selection is outside the PDF.")
            writer.add_page(reader.pages[idx])
        return _write(writer, safe_stem(filename)+"_organized.pdf")
    except AdvancedPdfError:
        raise
    except Exception as exc:
        raise AdvancedPdfError("Could not organize pages in this PDF.") from exc


def _overlay(width:float, height:float, text:str, position:str) -> bytes:
    """Text stays searchable for standard fonts; draw unicode via local Noto if available."""
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.lib.colors import Color
    from reportlab.lib.utils import ImageReader
    from PIL import Image, ImageDraw, ImageFont
    import os
    buffer=BytesIO(); c=Canvas(buffer, pagesize=(width,height), pageCompression=1)
    c.setFillColor(Color(.30,.30,.30,alpha=.20))
    size=max(12,min(26,width/18))
    y={"top":height-60,"center":height/2,"bottom":60}[position]
    if all(32 <= ord(ch) <= 126 for ch in text):
        c.setFont("Helvetica-Bold",size)
        c.drawCentredString(width/2,y,text)
    else:
        # Pillow's RAQM shaping support renders non-Latin labels consistently.
        script_font = ("NotoSansTamil-Regular.ttf" if any(0x0B80 <= ord(ch) <= 0x0BFF for ch in text)
                 else "NotoSansSinhala-Regular.ttf" if any(0x0D80 <= ord(ch) <= 0x0DFF for ch in text)
                 else "NotoSansArabic-Regular.ttf" if any(0x0600 <= ord(ch) <= 0x08FF for ch in text)
                 else "NotoSans-Regular.ttf")
        candidates=("/usr/share/fonts/truetype/noto/" + script_font,
                    "/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf",
                    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
                    "C:/Windows/Fonts/arial.ttf")
        font_path=next((p for p in candidates if os.path.isfile(p)),None)
        if not font_path:
            raise AdvancedPdfError("Unicode watermark font unavailable on this server.")
        font=ImageFont.truetype(font_path, 36)
        image=Image.new("RGBA",(max(1,int(width)),100),(255,255,255,0))
        draw=ImageDraw.Draw(image)
        draw.text((image.width/2,40),text,font=font,anchor="mm",fill=(80,80,80,55))
        c.drawImage(ImageReader(image),0,y-32,width=width,height=100,mask="auto")
        image.close()
    c.showPage();c.save()
    return buffer.getvalue()


def watermark_pdf(blob: bytes, filename: str, text: str, position: str) -> BytesIO:
    pdf_info(blob)
    text=str(text or "").strip()
    if not text or len(text)>60 or any(ord(ch)<32 for ch in text):
        raise AdvancedPdfError("Watermark text must have 1–60 characters and no line breaks.")
    if position not in ("top","center","bottom"):
        raise AdvancedPdfError("Choose Top, Center or Bottom watermark position.")
    from pypdf import PdfWriter, PdfReader
    try:
        reader=_pdf_reader(blob); writer=PdfWriter()
        for page in reader.pages:
            w,h=float(page.mediabox.width),float(page.mediabox.height)
            if w<55 or h<90:
                raise AdvancedPdfError("PDF page is too small for a watermark.")
            overlay=PdfReader(BytesIO(_overlay(w,h,text,position))).pages[0]
            page.merge_page(overlay,expand=False)
            writer.add_page(page)
        return _write(writer,safe_stem(filename)+"_watermarked.pdf")
    except AdvancedPdfError:
        raise
    except Exception as exc:
        raise AdvancedPdfError("Could not watermark this PDF.") from exc

# At most one expensive PDF job at a time on a small Render instance.
_worker_slot = threading.BoundedSemaphore(1)


def register_advanced_pdf_handlers(bot, open_pdf_tools, clear_pdf_state, show_main_menu):
    """Call before existing Office/PDF handlers to protect document routing."""
    from telebot import types
    from telebot.handler_backends import ContinueHandling

    states: dict[int, dict] = {}

    def cleanup_sessions():
        now = time.monotonic()
        for cid, session in tuple(states.items()):
            if now - session.get("last_activity", now) > SESSION_TTL_SECONDS:
                states.pop(cid, None)

    names = dict(zip(ADVANCED_BUTTONS, ("compress", "word", "images", "organize", "watermark")))

    def keyboard(*rows):
        kb = types.ReplyKeyboardMarkup(resize_keyboard=True, row_width=2)
        for row in rows:
            if isinstance(row, (list, tuple)):
                kb.row(*[types.KeyboardButton(x) for x in row])
            else:
                kb.add(types.KeyboardButton(row))
        return kb

    def nav_keyboard():
        return keyboard(["🔙 PDF Tools", "❌ Cancel"], "🏠 Main Menu")

    @bot.message_handler(func=lambda m: getattr(m, "text", None) in names, content_types=["text"])
    def choose_tool(message):
        chat_id = message.chat.id
        if getattr(message.chat, "type", "private") != "private":
            bot.send_message(chat_id, "🔒 Use PDF Tools in a private chat with the bot.")
            return
        cleanup_sessions()
        if chat_id not in states and len(states) >= MAX_SESSIONS:
            bot.send_message(chat_id, "⏳ PDF Tools are busy. Please try again shortly.")
            return
        clear_pdf_state(chat_id)
        tool = names[message.text]
        states[chat_id] = {"tool": tool, "stage": "upload", "last_activity": time.monotonic()}
        bot.send_message(chat_id, f"{message.text}\n\n📎 Upload one PDF (maximum 10 MB, 50 pages).\n"
                         "Encrypted PDFs are not supported.\nUse ❌ Cancel to exit.", reply_markup=nav_keyboard())

    @bot.message_handler(func=lambda m: m.chat.id in states and getattr(m, "text", "") in
                         {"🔙 PDF Tools", "❌ Cancel", "🏠 Main Menu"}, content_types=["text"])
    def cancel(message):
        states.pop(message.chat.id, None)
        if message.text == "🏠 Main Menu":
            show_main_menu(message.chat.id, getattr(message.from_user, "first_name", "Friend") or "Friend")
        else:
            open_pdf_tools(message.chat.id)

    @bot.message_handler(func=lambda m: m.chat.id in states and (
        str(getattr(m, "text", "") or "").startswith("/") or
        getattr(m, "text", "") in {"📄 PDF Tools", "🔗 Merge PDFs", "✂️ Extract Pages", "🖼 Images → PDF",
            "🖼 Photo/Image Crop", "📸 Scan Document", "📝 Text → PDF", "🔄 Rotate PDF",
            "🔐 Protect PDF", "📄 DOCX → PDF", "📊 PPTX → PDF", "🎙 Voice Notes", "🎙 Voice & PDF"}),
        content_types=["text"])
    def yield_to_other_tool(message):
        states.pop(message.chat.id, None)
        return ContinueHandling()

    def launch(chat_id: int, action, *, output_caption: str | None = None):
        state = states.get(chat_id)
        if not state or state.get("busy"):
            return
        if not _worker_slot.acquire(blocking=False):
            bot.send_message(chat_id, "⏳ Another PDF is being processed. Try again shortly.")
            return
        state["busy"] = True
        token = object()
        state["token"] = token
        data = state.pop("pdf", None)
        filename = state.get("filename", "document.pdf")

        def run():
            try:
                result = action(data, filename)
                if isinstance(result, tuple):
                    output, note = result
                else:
                    output, note = result, output_caption or "✅ Done."
                if states.get(chat_id, {}).get("token") is not token:
                    return
                try:
                    bot.send_document(chat_id, output, caption=note)
                finally:
                    output.close()
                states.pop(chat_id, None)
                open_pdf_tools(chat_id)
            except AdvancedPdfError as exc:
                if states.get(chat_id, {}).get("token") is token:
                    states.pop(chat_id, None)
                    bot.send_message(chat_id, f"❌ {exc}")
                    open_pdf_tools(chat_id)
            except Exception as exc:
                print("V7.17 PDF processing error:", type(exc).__name__)
                if states.get(chat_id, {}).get("token") is token:
                    states.pop(chat_id, None)
                    bot.send_message(chat_id, "❌ PDF processing failed. Please try a smaller document.")
                    open_pdf_tools(chat_id)
            finally:
                _worker_slot.release()
        threading.Thread(target=run, daemon=True, name="advanced-pdf-job").start()
        bot.send_message(chat_id, "⏳ Processing your PDF...")

    @bot.message_handler(func=lambda m: m.chat.id in states, content_types=["document"])
    def receive_document(message):
        chat_id = message.chat.id
        cleanup_sessions()
        state = states.get(chat_id)
        if not state or state.get("stage") != "upload" or state.get("busy"):
            bot.send_message(chat_id, "⏳ Complete or cancel the current PDF task first.")
            return
        file = message.document
        filename = str(getattr(file, "file_name", "") or "")
        if Path(filename.replace("\\", "/")).suffix.lower() != ".pdf":
            bot.send_message(chat_id, "❌ Upload a .pdf file for this tool.")
            return
        if (getattr(file, "file_size", None) or 0) > MAX_INPUT_BYTES:
            bot.send_message(chat_id, "❌ PDF must be 10 MB or smaller.")
            return
        try:
            info = bot.get_file(file.file_id)
            if (getattr(info, "file_size", None) or 0) > MAX_INPUT_BYTES:
                raise AdvancedPdfError("PDF must be 10 MB or smaller.")
            blob = bot.download_file(info.file_path)
            pages = pdf_info(blob)
            state.update(pdf=blob, filename=filename, total=pages)
        except AdvancedPdfError as exc:
            bot.send_message(chat_id, f"❌ {exc}")
            return
        except Exception as exc:
            print("V7.17 PDF download error:", type(exc).__name__)
            bot.send_message(chat_id, "❌ Download failed. Try uploading the PDF again.")
            return
        state["last_activity"] = time.monotonic()
        tool = state["tool"]
        if tool == "word":
            launch(chat_id, pdf_to_word, output_caption="✅ Editable Word file created from selectable PDF text. Check complex layouts.")
        elif tool == "compress":
            state["stage"] = "level"
            bot.send_message(chat_id, "🗜 Choose compression: Light (lossless), Balanced or Strong (recompress images).",
                             reply_markup=keyboard(["🟢 Light", "🟡 Balanced", "🔴 Strong"], "🔙 PDF Tools"))
        elif tool == "images":
            state["stage"] = "format"
            bot.send_message(chat_id, "🖼 Choose image format (maximum 20 pages, ZIP output):",
                             reply_markup=keyboard(["🖼 PNG", "📷 JPG"], "🔙 PDF Tools"))
        elif tool == "organize":
            state["stage"] = "pages"
            bot.send_message(chat_id, f"🔀 PDF has {pages} pages. Send the final page order, for example: 3,1,2-5.\n"
                             "Omit pages to delete them; repeat a page to duplicate it.\n"
                             "Send All to keep every page in the same order.", reply_markup=nav_keyboard())
        else:
            state["stage"] = "watermark_text"
            bot.send_message(chat_id, "💧 Type your watermark text (1–60 characters). Example: DRAFT",
                             reply_markup=nav_keyboard())

    @bot.message_handler(func=lambda m: m.chat.id in states, content_types=["text"])
    def receive_option(message):
        chat_id = message.chat.id
        cleanup_sessions()
        state = states.get(chat_id)
        if not state:
            bot.send_message(chat_id, "⌛ PDF session expired. Open PDF Tools and start again.")
            return
        state["last_activity"] = time.monotonic()
        if state.get("busy"):
            bot.send_message(chat_id, "⏳ PDF processing is in progress. Tap Cancel to exit.")
            return
        text = (message.text or "").strip()
        stage = state.get("stage")
        try:
            if stage == "upload":
                bot.send_message(chat_id, "📎 Please upload your PDF as a Telegram document.")
            elif stage == "level":
                if text not in COMPRESS_LEVELS:
                    raise AdvancedPdfError("Choose Light, Balanced or Strong using the buttons.")
                launch(chat_id, lambda blob, filename: compress_pdf(blob, filename, text))
            elif stage == "format":
                if text not in IMAGE_FORMATS:
                    raise AdvancedPdfError("Choose PNG or JPG.")
                state["format"] = IMAGE_FORMATS[text]
                state["stage"] = "pages"
                bot.send_message(chat_id, f"📑 {state['total']} pages. Send All or a list like 1,3-5 (max 20 pages).",
                                 reply_markup=keyboard("✅ All Pages", "🔙 PDF Tools"))
            elif stage == "pages":
                is_images = state["tool"] == "images"
                chosen = parse_pages(text, state["total"], max_pages=MAX_IMAGE_PAGES if is_images else MAX_OUTPUT_PAGES,
                                     preserve_duplicates=not is_images)
                if is_images:
                    fmt = state["format"]
                    launch(chat_id, lambda blob, filename: pdf_to_images(blob, filename, fmt, chosen),
                           output_caption=f"✅ Exported {len(chosen)} pages as {fmt.upper()} images (ZIP).")
                else:
                    launch(chat_id, lambda blob, filename: organize_pdf(blob, filename, chosen),
                           output_caption=f"✅ Organized PDF with {len(chosen)} pages.")
            elif stage == "watermark_text":
                if not text or len(text) > 60 or any(ord(c) < 32 for c in text):
                    raise AdvancedPdfError("Enter 1–60 characters of watermark text.")
                state["watermark"] = text
                state["stage"] = "position"
                bot.send_message(chat_id, "💧 Choose watermark location (light, transparent text):",
                                 reply_markup=keyboard(["⬆️ Top", "🎯 Center", "⬇️ Bottom"], "🔙 PDF Tools"))
            elif stage == "position":
                if text not in WATERMARK_PLACEMENTS:
                    raise AdvancedPdfError("Choose Top, Center or Bottom using the buttons.")
                position = WATERMARK_PLACEMENTS[text]
                watermark = state["watermark"]
                launch(chat_id, lambda blob, filename: watermark_pdf(blob, filename, watermark, position),
                       output_caption="✅ Watermark added to every page.")
            else:
                raise AdvancedPdfError("Choose a tool from PDF Tools.")
        except AdvancedPdfError as exc:
            bot.send_message(chat_id, f"❌ {exc}")

    return states
