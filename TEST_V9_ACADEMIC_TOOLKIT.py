"""Run: python -m unittest TEST_V9_ACADEMIC_TOOLKIT -q

Offline tests. No Render, Telegram, OCR.space, Crossref or Supabase network calls.
"""
import io, os, sys, types, tempfile, unittest
from pathlib import Path
from unittest.mock import patch, MagicMock

from v8_engine import StudyInputError
from v9_engine import (math_solution, assignment_report, syllabus_lines,
    formula_lines, formula_docx, prepare_deck, document_diff, mindmap_topics,
    mindmap_png, presentation_pptx, parse_timetable, diagram_labels)
from INSTALL_V9_ACADEMIC_TOOLKIT import patched


class EngineTests(unittest.TestCase):
    def test_math_linear(self):self.assertIn('x = 4',math_solution('2*x+3=11'))
    def test_math_quadratic(self):
        result=math_solution('x^2-4=0'); self.assertIn('x = -2',result);self.assertIn('x = 2',result)
    def test_math_division(self):self.assertIn('x = 10',math_solution('x/2=5'))
    def test_math_explicit_multiplication(self):
        with self.assertRaises(StudyInputError):math_solution('2x+3=11')
    def test_math_attack(self):
        for equation in ('__import__("os")=0','x**999=0','x+z=2','x=1=2','x^2^2^2=9'):
            with self.subTest(equation=equation):
                with self.assertRaises(StudyInputError):math_solution(equation)
    def test_assignment(self):
        data=('Introduction\nStudy the methods.\nConclusion\nUseful result.\nReferences\nSmith (2023).').encode()
        result=assignment_report('work.txt',data,min_words=6,max_words=100)
        self.assertIn('✅ Introduction',result)
        self.assertIn('Word limit: ✅',result)
    def test_syllabus(self):self.assertEqual(syllabus_lines('1. Chemistry\n2. Physics\n3. Physics'),['Chemistry','Physics'])
    def test_syllabus_too_short(self):
        with self.assertRaises(StudyInputError):syllabus_lines('A\nB')
    def test_formula_lines(self):
        result=formula_lines('Speed = Distance / Time\nThe weather is good.\nx = 2')
        self.assertEqual(len(result),2)
    def test_formula_docx(self):
        from docx import Document
        doc=formula_docx(['v = s / t'])
        self.assertIn('v = s / t','\n'.join(p.text for p in Document(doc).paragraphs))
    def test_anki_tsv(self):
        raw=prepare_deck([('Q1','A1'),('=2+2','Four')]).read().decode('utf-8-sig')
        self.assertIn('Front\tBack',raw);self.assertIn("'=2+2",raw)
    def test_document_diff(self):
        self.assertIn('+three',document_diff('one\ntwo','one\nthree'))
    def test_document_same(self):self.assertIn('No line-level changes',document_diff('Same','Same'))
    def test_mindmap_topics(self):self.assertIn('Biology',mindmap_topics('Biology\nPhotosynthesis\nAnatomy'))
    def test_mindmap_png(self):
        out=mindmap_png('Science',['Physics','Biology'])
        self.assertEqual(out.read(8),b'\x89PNG\r\n\x1a\n')
    def test_pptx(self):
        from pptx import Presentation
        out,n=presentation_pptx('\n'.join('Topic '+str(i) for i in range(25)))
        self.assertEqual(len(Presentation(out).slides),n)
        self.assertGreater(n,2)
    def test_pptx_empty(self):
        with self.assertRaises(StudyInputError):presentation_pptx('')
    def test_timetable(self):
        rows=parse_timetable('Mon 09:00-10:30 Math\nThu 11:00-12:00 English\nBAD LINE')
        self.assertEqual([v[3] for v in rows],['Math','English'])
    def test_timetable_invalid(self):self.assertFalse(parse_timetable('Mon 14:00-12:00 Bad'))
    def test_diagram_honesty(self):
        result=diagram_labels('Nucleus\nCell wall')
        self.assertIn('does NOT infer arrows',result)
    def test_diagram_empty(self):
        with self.assertRaises(StudyInputError):diagram_labels('a')


class InstallerTests(unittest.TestCase):
    def test_main_insertion(self):
        source='# V8.0 Study Lab — opt-in and registered before stateful legacy workflows.\n'
        result=patched('main.py',source)
        self.assertIn('register_v9_handlers',result)
        self.assertEqual(result,patched('main.py',result))
    def test_ui(self):
        result=patched('ui_v57.py','ALL_UI_BUTTONS = (')
        self.assertIn('V9_ENABLED',result)
        self.assertEqual(result,patched('ui_v57.py',result))
    def test_navigation(self):
        result=patched('navigation.py','GLOBAL_BUTTONS.add("📝 Text → PDF")')
        self.assertIn('🧰 V9 Academic Toolkit',result)
    def test_requires_v8_anchor(self):
        with self.assertRaises(ValueError):patched('main.py','old source')
    def test_requirements_idempotent(self):
        a=patched('requirements.txt','Flask>=3\n')
        self.assertIn('sympy',a)
        self.assertEqual(a,patched('requirements.txt',a))
    def test_actual_project_files(self):
        root=Path(__file__).resolve().parent
        main=root/'main.py'
        if not main.exists():self.skipTest('Install into existing V8 project to verify anchors')
        for name in ('main.py','ui_v57.py','navigation.py','requirements.txt'):
            text=(root/name).read_text(encoding='utf-8')
            result=patched(name,text)
            self.assertEqual(patched(name,result),result)


class BotTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        try:import telebot
        except ImportError:
            mod=types.ModuleType('telebot')
            class Keyboard:
                def __init__(self,*args,**kw):self.buttons=[]
                def row(self,*args):self.buttons.append(args)
            mod.types=types.SimpleNamespace(ReplyKeyboardMarkup=Keyboard,KeyboardButton=lambda x:x)
            back=types.ModuleType('telebot.handler_backends');back.ContinueHandling=type('ContinueHandling',(),{})
            sys.modules['telebot']=mod;sys.modules['telebot.handler_backends']=back
        import v9_bot
        cls.mod=v9_bot
    def setUp(self):
        self.env=patch.dict(os.environ,{'V8_ENABLED':'true','V8_ALLOWED_USER_IDS':'123','V9_ENABLED':'true'})
        self.env.start()
        class Bot:
            def __init__(self):self.handlers={};self.sent=[]
            def message_handler(self,**kwargs):
                def register(f):self.handlers[f.__name__]=f;return f
                return register
            def send_message(self,chat_id,text,reply_markup=None):self.sent.append((chat_id,text,reply_markup))
            def send_photo(self,chat_id,media,caption=''):self.sent.append((chat_id,caption,None))
            def send_document(self,chat_id,media,caption=''):self.sent.append((chat_id,caption,None))
        self.bot=Bot()
        self.mod._states.clear()
        self.mod.register_v9_handlers(self.bot,lambda x:None)
    def tearDown(self):self.env.stop();self.mod._states.clear()
    def msg(self, text=None, uid=123, chat_type='private'):
        return types.SimpleNamespace(text=text,from_user=types.SimpleNamespace(id=uid),
                                      chat=types.SimpleNamespace(id=uid,type=chat_type))
    def test_hub_shows_14_tools(self):
        self.bot.handlers['command_hub'](self.msg('/v9'))
        self.assertIn('14 student tools',self.bot.sent[-1][1])
    def test_unauthorized(self):
        self.bot.handlers['command_hub'](self.msg('/v9',456))
        self.assertIn('approved',self.bot.sent[-1][1])
    def test_group_rejected(self):
        self.bot.handlers['command_hub'](self.msg('/v9',chat_type='group'))
        self.assertIn('private',self.bot.sent[-1][1])
    def test_math_workflow(self):
        self.bot.handlers['menu_feature'](self.msg('📸 Math Solver'))
        self.bot.handlers['text_handler'](self.msg('2*x+3=11'))
        self.assertIn('x = 4',self.bot.sent[-1][1])
    def test_requires_ocr_consent(self):
        self.mod._set(123,{'kind':'math','stage':'ocr_consent','data':b'test'})
        with patch.object(self.mod,'_read_ocr') as ocr:
            self.bot.handlers['text_handler'](self.msg('NO'))
            ocr.assert_not_called()
        self.assertIn('cancelled',self.bot.sent[-1][1])
    def test_ocr_consented(self):
        self.mod._set(123,{'kind':'math','stage':'ocr_consent','data':b'test'})
        with patch.object(self.mod,'_read_ocr',return_value='2*x+3=11') as ocr:
            self.bot.handlers['text_handler'](self.msg('I AGREE OCR'))
            ocr.assert_called_once()
        self.assertIn('x = 4',self.bot.sent[-1][1])
    def test_cancel(self):
        self.mod._set(123,{'kind':'math','stage':'text'})
        self.bot.handlers['cancel'](self.msg('/v9cancel'))
        self.assertIsNone(self.mod._get(123))
    def test_tool_buttons(self):
        self.assertEqual(len(self.mod.OPTIONS),14)
        for label,kind in self.mod.OPTIONS.items():self.assertEqual(self.mod.TOOLS[kind],label)
    def test_flags(self):
        os.environ['V9_ENABLED']='false'
        self.assertFalse(self.mod.allowed(123))
    def test_doi_validation(self):
        with self.assertRaises(StudyInputError):self.mod.verify_reference('http://127.0.0.1/secret')
    def test_doi_mock(self):
        class Reply:
            status_code=200
            def raise_for_status(self):pass
            def json(self):return {'message':{'title':['Learning Science'],'published':{'date-parts':[[2023]]},'DOI':'10.1234/a'}}
        with patch('requests.get',return_value=Reply()) as get:
            result=self.mod.verify_reference('10.1234/a')
            self.assertIn('Learning Science',result)
            self.assertEqual(get.call_args.args[0].split('/')[2],'api.crossref.org')

if __name__=='__main__':unittest.main()
