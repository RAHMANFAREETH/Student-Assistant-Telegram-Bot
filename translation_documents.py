from io import BytesIO
from pathlib import Path

from PyPDF2 import PdfReader
from docx import Document


SUPPORTED_DOCUMENT_EXTENSIONS = {
    ".txt",
    ".pdf",
    ".docx",
}


def extract_text_from_document(file_bytes, filename):
    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_DOCUMENT_EXTENSIONS:
        raise ValueError(
            "Unsupported document type. Please send TXT, PDF, or DOCX."
        )

    if extension == ".txt":
        return _extract_txt(file_bytes)

    if extension == ".pdf":
        return _extract_pdf(file_bytes)

    if extension == ".docx":
        return _extract_docx(file_bytes)

    raise ValueError("Unsupported document type.")


def _extract_txt(file_bytes):
    for encoding in ("utf-8", "utf-8-sig", "utf-16"):
        try:
            text = file_bytes.decode(encoding)
            if text.strip():
                return text
        except UnicodeDecodeError:
            continue

    return file_bytes.decode("utf-8", errors="replace")


def _extract_pdf(file_bytes):
    reader = PdfReader(BytesIO(file_bytes))
    pages = []

    for page in reader.pages:
        text = page.extract_text() or ""

        if text.strip():
            pages.append(text.strip())

    return "\n\n".join(pages).strip()


def _extract_docx(file_bytes):
    document = Document(BytesIO(file_bytes))
    blocks = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            blocks.append(paragraph.text.strip())

    for table in document.tables:
        for row in table.rows:
            values = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]

            if values:
                blocks.append(" | ".join(values))

    return "\n\n".join(blocks).strip()


def build_translated_text_file(translated_text, original_filename):
    original = Path(original_filename)
    output_name = f"{original.stem}_translated.txt"

    output = BytesIO(
        translated_text.encode("utf-8")
    )

    output.name = output_name
    output.seek(0)

    return output, output_name
