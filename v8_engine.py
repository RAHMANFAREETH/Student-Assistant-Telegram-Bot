"""V8.0 offline study utilities. Input is untrusted and bounded; no external APIs.

This module deliberately does not claim AI-generated inference. Existing AI workflows
remain available in the bot and are never called without their own authorization.
"""
from __future__ import annotations

from collections import Counter
from datetime import date, datetime, timezone, timedelta
from io import BytesIO
import re
import unicodedata

MAX_FILE_BYTES = 8 * 1024 * 1024
MAX_PAGES = 35
MAX_TEXT = 100_000
WORD = re.compile(r"[^\W_]{2,}", re.UNICODE)
STOP = {"the", "and", "for", "this", "with", "that", "from", "what", "when", "were", "which", "will", "are", "have", "your", "into", "about", "define", "explain", "write", "discuss", "state", "list", "answer", "question", "each", "these"}

class StudyInputError(ValueError):
    pass


def tokens(text: str) -> list[str]:
    return [t.lower() for t in WORD.findall(text or "") if t.lower() not in STOP and len(t) > 2]


def _check_size(data):
    if not isinstance(data, bytes) or not 0 < len(data) <= MAX_FILE_BYTES:
        raise StudyInputError("Upload a file up to 8 MB.")


def extract_document(filename: str, data: bytes) -> list[tuple[int, str]]:
    """Return [(page, text)] with page numbers for PDF; DOCX/TXT = page 1.

    Read textual content only. Scans/images require existing OCR tools.
    """
    _check_size(data)
    low = (filename or "").lower()
    output = []
    try:
        if low.endswith(".pdf"):
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(data), strict=False)
            if reader.is_encrypted:
                raise StudyInputError("Password-protected PDFs are not supported.")
            if not 0 < len(reader.pages) <= MAX_PAGES:
                raise StudyInputError(f"Use a PDF with 1–{MAX_PAGES} pages.")
            for i, page in enumerate(reader.pages, 1):
                text = page.extract_text() or ""
                if text.strip():
                    output.append((i, text[:MAX_TEXT]))
        elif low.endswith(".docx"):
            from docx import Document
            doc = Document(BytesIO(data))
            text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
            # Include table cells where relevant, without duplicating paragraphs.
            for table in doc.tables[:15]:
                for row in table.rows[:100]:
                    text += "\n" + " | ".join(cell.text[:500] for cell in row.cells)
            output = [(1, text[:MAX_TEXT])]
        elif low.endswith(".txt"):
            try:
                text = data.decode("utf-8-sig")
            except UnicodeDecodeError as exc:
                raise StudyInputError("Text files must use UTF-8 encoding.") from exc
            output = [(1, text[:MAX_TEXT])]
        else:
            raise StudyInputError("Supported documents: .pdf, .docx, .txt.")
    except StudyInputError:
        raise
    except Exception as exc:
        raise StudyInputError("Could not read this file. Try another document.") from exc
    remaining = MAX_TEXT
    clean = []
    for page, text in output:
        text = text.replace("\x00", "").strip()[:remaining]
        if text:
            clean.append((page, text))
            remaining -= len(text)
        if remaining <= 0:
            break
    if sum(len(text) for _, text in clean) < 10:
        raise StudyInputError("No selectable text found. Use the existing OCR/Scanner tools for scanned pages.")
    return clean


def answer_from_document(pages: list[tuple[int, str]], question: str, limit: int = 4) -> str:
    """Grounded extractive Q&A; never fabricates an answer or calls AI."""
    if not 3 <= len(question.strip()) <= 350:
        raise StudyInputError("Write a question of 3–350 characters.")
    q = set(tokens(question))
    if not q:
        raise StudyInputError("Include a specific topic or keyword in your question.")
    ranked = []
    for page_num, text in pages:
        for part in re.split(r"(?<=[.!?؟।])\s+|\n{2,}", text):
            part = " ".join(part.split())
            if not 24 <= len(part) <= 900:
                continue
            score = len(set(tokens(part)) & q)
            if score:
                ranked.append((score, page_num, part))
    ranked.sort(key=lambda x: (-x[0], x[1], len(x[2])))
    if not ranked:
        return "I couldn't find matching passages in the uploaded text. Try different keywords."
    distinct = set(); lines = []
    for _, page, excerpt in ranked:
        key = excerpt[:120].lower()
        if key in distinct:
            continue
        distinct.add(key)
        lines.append(f"[Page {page}] {excerpt[:520]}")
        if len(lines) >= limit:
            break
    return "🔎 Matching passages (not a generated conclusion):\n\n" + "\n\n".join(lines) + "\n\nCheck the original document before citing."


def analyze_past_paper(pages: list[tuple[int, str]]) -> dict:
    """Offline extraction of numbered questions and recurring terms; heuristic only."""
    text = "\n".join(t for _, t in pages)
    questions = []
    for line in text.splitlines():
        clean = " ".join(line.split())
        if len(clean) < 12 or len(clean) > 400:
            continue
        if re.match(r"(?i)^(?:q(?:uestion)?\s*\d+[.):\-]?|\d{1,2}[.):]|\([a-z]\))\s+", clean) or clean.endswith("?"):
            questions.append(clean)
    counts = Counter(tokens(" ".join(questions) if questions else text))
    return {"questions": list(dict.fromkeys(questions))[:25], "topics": counts.most_common(10), "total_detected": len(questions)}


def flashcards_from_text(pages: list[tuple[int,str]], limit=10) -> list[tuple[str,str]]:
    """Produce deterministic definition-based cards only where a source supports them."""
    result = []; seen = set()
    for page, text in pages:
        for raw in text.splitlines():
            line = " ".join(raw.split())
            if not 12 <= len(line) <= 450: continue
            match = re.match(r"^(.{3,65}?)(?::|\s+[–—-]\s+|\s+is\s+(?:a|an|the)\s+)(.{8,330})$", line, re.IGNORECASE)
            if not match: continue
            term, definition = (x.strip(' .:-') for x in match.groups())
            if len(term.split()) > 9 or len(definition) < 10: continue
            if term.lower() in seen: continue
            seen.add(term.lower()); result.append((f"What is {term}? (p. {page})", definition))
            if len(result) >= limit: return result
    return result


def practice_quiz(cards: list[tuple[str,str]], size=5):
    """Create MCQs only when 4+ distinct definitions are available."""
    answers = list(dict.fromkeys(a for _, a in cards))
    if len(answers) < 4:
        raise StudyInputError("At least four distinct flashcard definitions are needed for multiple-choice questions.")
    quiz=[]
    for idx,(question,answer) in enumerate(cards[:size]):
        distractors = [a for a in answers if a!=answer][:3]
        options = [answer]+distractors
        rot = idx % 4
        options = options[rot:]+options[:rot]
        quiz.append((question, options, "ABCD"[options.index(answer)]))
    return quiz


def make_exam_plan(exams: list[dict], daily_minutes: int = 60, today: date | None = None) -> list[str]:
    if not 15 <= daily_minutes <= 360:
        raise StudyInputError("Study time must be 15–360 minutes per day.")
    today = today or date.today()
    parsed = []
    for row in exams:
        try:
            timestamp = datetime.fromisoformat(str(row["exam_at"]).replace("Z", "+00:00"))
            when = timestamp.date()
            if when >= today:
                parsed.append((when, str(row.get("title") or "Exam")[:80]))
        except (KeyError, TypeError, ValueError): continue
    if not parsed:
        return ["No upcoming exams found. Add exams through 📝 Exams first."]
    parsed.sort()
    plan=[]
    for d in range(0, min(7, (parsed[-1][0]-today).days+1)):
        day = today + timedelta(days=d)
        remaining = [(exam,title) for exam,title in parsed if exam >= day]
        if not remaining: break
        # Allocate more minutes to nearer exams; avoid assuming actual student proficiency.
        weights = [1/max(1,(exam-day).days+1) for exam,_ in remaining]
        chosen = sorted(zip(remaining,weights),key=lambda x:-x[1])[:3]
        total = sum(weight for _,weight in chosen)
        sessions = [f"{title}: {round(daily_minutes*weight/total/5)*5 or 5} min" for (exam,title),weight in chosen]
        plan.append(f"{day.isoformat()}: " + " • ".join(sessions))
    return plan


def summarize_lecture(transcript: str) -> str:
    if len(transcript.strip()) < 40:
        raise StudyInputError("No usable transcript. Transcribe a Voice Note first.")
    words = Counter(tokens(transcript))
    sentences = [" ".join(s.split()) for s in re.split(r"(?<=[.!?؟])\s+|\n+", transcript) if len(s.strip()) > 20]
    ranked = sorted(enumerate(sentences), key=lambda kv: -sum(words[t] for t in set(tokens(kv[1]))))
    selected = sorted(ranked[:min(5,len(ranked))])
    key_terms = ", ".join(term for term,_ in words.most_common(8))
    return "🎙 Key passages from the transcript (automated, verify accuracy):\n" + "\n".join(f"• {s[:320]}" for _,s in selected) + "\n\nKeywords: " + (key_terms or "not found")


def citation(author: str, year: str, title: str, source: str, style: str="APA") -> str:
    fields = [unicodedata.normalize("NFKC", str(s)).strip() for s in (author,year,title,source)]
    if any(not x or len(x)>250 for x in fields):
        raise StudyInputError("Use nonempty author, year, title and source (max 250 chars each).")
    author,year,title,source=fields
    if not re.fullmatch(r"\d{4}|n\.d\.", year):
        raise StudyInputError("Year must be four digits or n.d.")
    if style.upper()=="APA": return f"{author} ({year}). {title}. {source}."
    if style.upper()=="HARVARD": return f"{author} ({year}) '{title}', {source}."
    if style.upper()=="MLA": return f"{author}. \"{title}.\" {source}, {year}."
    raise StudyInputError("Use APA, Harvard or MLA.")
